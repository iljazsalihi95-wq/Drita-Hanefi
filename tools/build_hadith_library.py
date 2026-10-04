#!/usr/bin/env python3
"""Build the offline Hadith catalogue from the owner-provided 4872 LIVE workbook."""

import argparse
import html
import json
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


COLLECTIONS = {
    "abudawud": "Sunen Ebu Davud",
    "adab": "El-Edeb el-Mufred",
    "ahmad": "Musned Ahmed",
    "bukhari": "Sahih el-Buhari",
    "bulugh": "Bulug el-Meram",
    "forty": "Dyzet Hadithet e Neveviut",
    "hisn": "Hisn el-Muslim",
    "ibnmajah": "Sunen Ibn Maxhe",
    "malik": "Muvatta Malik",
    "mishkat": "Mishkat el-Mesabih",
    "muslim": "Sahih Muslim",
    "nasai": "Sunen en-Nesai",
    "riyadussalihin": "Rijad es-Salihin",
    "shamail": "Esh-Shemail el-Muhammedije",
    "tirmidhi": "Xhami et-Tirmidhi",
}

CURATED_SHEETS = [
    "BUHARI_V4", "SAHIH_MUSLIM", "SUNEN_EBU_DAVUD", "JAMI_TIRMIDHI",
    "SUNEN_NESAI", "SUNEN_IBN_MAXHE",
]


def clean(value):
    if value is None:
        return ""
    return str(value).strip()


def clean_number(value):
    text = clean(value)
    return text[:-2] if re.fullmatch(r"-?\d+\.0", text) else text


def strip_html(value):
    text = re.sub(r"<[^>]+>", "", clean(value))
    return html.unescape(text).replace("\u200f", "").strip()


def rows(ws):
    iterator = ws.iter_rows(values_only=True)
    header = [clean(x) for x in next(iterator)]
    for values in iterator:
        row = {header[i]: values[i] for i in range(min(len(header), len(values))) if header[i]}
        if any(clean(v) for v in row.values()):
            yield row


def collection_id(name):
    value = clean(name).lower()
    aliases = {
        "sahih el-buhari": "bukhari", "sahih bukhari": "bukhari",
        "sahih muslim": "muslim", "sunan ebu davud": "abudawud",
        "sunen ebu davud": "abudawud", "xhami et-tirmidhi": "tirmidhi",
        "jami tirmidhi": "tirmidhi", "sunan en-nesai": "nasai",
        "sunen en-nesai": "nasai", "sunan ibn maxhe": "ibnmajah",
        "sunen ibn maxhe": "ibnmajah",
    }
    return aliases.get(value, value)


def is_placeholder(value):
    return bool(re.search(r"PËR (?:TRANSKRIPTIM|VERIFIKIM)|DO TË|PLACEHOLDER|DEMO", clean(value), re.I))


def publishable(value):
    text = clean(value)
    return "" if is_placeholder(text) else text


def curated_index(workbook):
    result = {}
    for sheet in CURATED_SHEETS:
        if sheet not in workbook.sheetnames:
            continue
        for row in rows(workbook[sheet]):
            cid = collection_id(row.get("Koleksioni"))
            number = clean_number(row.get("Nr. Hadithit"))
            translation = clean(row.get("Përkthimi Shqip"))
            if not cid or not number or not translation or is_placeholder(translation):
                continue
            result[(cid, number)] = row
    return result


def compact(raw, curated):
    cid = clean(raw.get("Collection")).lower()
    number = clean_number(raw.get("HadithNumber"))
    label = COLLECTIONS.get(cid, clean(raw.get("CollectionSq")) or cid)
    editorial = curated.get((cid, number), {})
    arabic = strip_html(raw.get("HadithTextAr"))
    editorial_arabic = clean(editorial.get("Teksti Arab"))
    if editorial_arabic and not is_placeholder(editorial_arabic):
        arabic = strip_html(editorial_arabic)
    translation = clean(editorial.get("Përkthimi Shqip")) or clean(raw.get("PerkthimiShqip"))
    title = clean(editorial.get("Tema")) or clean(raw.get("TitulliSq")) or f"{label} – Hadithi {number}"
    book = clean(editorial.get("Kitab / Libri")) or clean(raw.get("BookTitleSq"))
    if not book and clean(raw.get("BookNumber")):
        book = f"Libri nr. {clean_number(raw.get('BookNumber'))}"
    chapter = clean(editorial.get("Bab / Kapitulli")) or clean(raw.get("ChapterTitleSq"))
    if not chapter and clean(raw.get("ChapterId")):
        chapter = f"Kapitulli nr. {clean_number(raw.get('ChapterId'))}"
    grade = clean(editorial.get("Gradimi")) or clean(raw.get("HadithGradedAr"))
    grader = clean(editorial.get("Graduesi")) or clean(raw.get("HadithGradedByAr"))
    record = {
        "id": f"{cid}-{clean_number(raw.get('HadithId')) or number}",
        "collection": cid,
        "collectionLabel": label,
        "number": number,
        "numberInBook": clean_number(raw.get("HadithNumberInBook")),
        "bookNumber": clean_number(raw.get("BookNumber")),
        "chapterId": clean_number(raw.get("ChapterId")),
        "title": title,
        "topic": clean(editorial.get("Tema")) or clean(raw.get("TemaSq")),
        "subtopic": clean(editorial.get("Nëntema")) or clean(raw.get("NentemaSq")),
        "book": book,
        "chapter": chapter,
        "arabic": arabic,
        "transliteration": publishable(editorial.get("Transkriptimi")) or publishable(raw.get("Transkriptimi")),
        "albanian": translation,
        "chain": clean(editorial.get("Senedi")) or clean(raw.get("SenediShqip")),
        "companion": clean(editorial.get("Sahabiu")) or clean(raw.get("Sahabiu")),
        "grade": grade,
        "grader": grader,
        "lesson": clean(editorial.get("Hukmi / Mësimet")) or clean(raw.get("HukmiNxjerre")),
        "commentary": clean(editorial.get("Koment / Sherh")),
        "hanafiCommentary": publishable(editorial.get("Koment Hanefi")) or publishable(raw.get("KomentHanefi")),
        "hanafiSources": publishable(editorial.get("Burime Hanefi")) or publishable(raw.get("BurimeHanefi")),
        "reference": clean(editorial.get("Burimi origjinal")) or f"{label}, hadithi {number}",
        "sourceUrl": clean(editorial.get("URL Burimi")),
        "status": publishable(editorial.get("Status Verifikimi")) or publishable(raw.get("StatusVerifikimi")),
    }
    return {key: value for key, value in record.items() if value not in ("", None)}


def curated_record(cid, number, row):
    arabic = clean(row.get("Teksti Arab"))
    translation = clean(row.get("Përkthimi Shqip"))
    if not arabic or is_placeholder(arabic) or not translation or is_placeholder(translation):
        return None
    label = COLLECTIONS.get(cid, clean(row.get("Koleksioni")))
    record = {
        "id": f"{cid}-sq-{number}", "collection": cid, "collectionLabel": label,
        "number": number, "title": clean(row.get("Tema")) or f"{label} – Hadithi {number}",
        "topic": clean(row.get("Tema")), "subtopic": clean(row.get("Nëntema")),
        "book": clean(row.get("Kitab / Libri")), "chapter": clean(row.get("Bab / Kapitulli")),
        "arabic": strip_html(arabic), "transliteration": publishable(row.get("Transkriptimi")),
        "albanian": translation, "chain": clean(row.get("Senedi")),
        "companion": clean(row.get("Sahabiu")), "grade": clean(row.get("Gradimi")),
        "grader": clean(row.get("Graduesi")), "lesson": clean(row.get("Hukmi / Mësimet")),
        "commentary": publishable(row.get("Koment / Sherh")), "hanafiCommentary": publishable(row.get("Koment Hanefi")),
        "hanafiSources": publishable(row.get("Burime Hanefi")), "reference": clean(row.get("Burimi origjinal")),
        "sourceUrl": clean(row.get("URL Burimi")), "status": publishable(row.get("Status Verifikimi")),
        "editorialExtra": True,
    }
    return {key: value for key, value in record.items() if value not in ("", None)}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook")
    parser.add_argument("output")
    args = parser.parse_args()
    workbook = load_workbook(args.workbook, read_only=True, data_only=True)
    editorial = curated_index(workbook)
    output = Path(args.output)
    output.mkdir(parents=True, exist_ok=True)
    grouped = {key: [] for key in COLLECTIONS}
    source = workbook["Drita Hanefi – Hadith 4872 LIVE"]
    source_keys = set()
    for row in rows(source):
        record = compact(row, editorial)
        grouped.setdefault(record["collection"], []).append(record)
        source_keys.add((record["collection"], record["number"]))
    editorial_extras = 0
    for (cid, number), row in editorial.items():
        if (cid, number) in source_keys:
            continue
        record = curated_record(cid, number, row)
        if record:
            grouped.setdefault(cid, []).append(record)
            editorial_extras += 1
    collections = []
    total = translated = graded = 0
    for cid, records in grouped.items():
        if not records:
            continue
        records.sort(key=lambda x: (int(x.get("number", 0)) if str(x.get("number", "")).isdigit() else 10**9, x["id"]))
        file_name = f"{cid}.json"
        (output / file_name).write_text(json.dumps({"collection": cid, "rows": records}, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
        counts = Counter({"total": len(records), "translated": sum(bool(x.get("albanian")) for x in records), "graded": sum(bool(x.get("grade")) for x in records)})
        total += counts["total"]
        translated += counts["translated"]
        graded += counts["graded"]
        collections.append({"id": cid, "title": COLLECTIONS.get(cid, cid), "file": file_name, **counts})
    manifest = {
        "source": {"title": "Drita Hanefi – Hadith 4872 LIVE", "spreadsheetId": "12i8lAsSDHumk8NiYajhqVth34DRtWnjoba43v5HZZWE", "sheet": "Drita Hanefi – Hadith 4872 LIVE"},
        "baseTotal": 4872, "total": total, "translated": translated, "graded": graded,
        "editorialExtras": editorial_extras,
        "collections": sorted(collections, key=lambda x: x["title"]),
    }
    (output / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"total": total, "translated": translated, "graded": graded, "collections": len(collections), "editorial": len(editorial)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
